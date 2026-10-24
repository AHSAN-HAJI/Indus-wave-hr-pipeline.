from pptx import Presentation

prs = Presentation()

# Slide 1: Title Slide
slide_layout = prs.slide_layouts[0]
slide = prs.slides.add_slide(slide_layout)
slide.shapes.title.text = "Indus Weave HR Pipeline"
slide.placeholders[1].text = "CFO Data Warehouse & Overtime Cost Analysis\nPrepared for Hackathon Demo"

# Slide 2: The Problem
slide_layout = prs.slide_layouts[1]
slide = prs.slides.add_slide(slide_layout)
slide.shapes.title.text = "The Problem: Rising Overtime Bills"
tf = slide.placeholders[1].text_frame
tf.text = "• Overtime bill has climbed continuously for 6 months."
tf.add_paragraph().text = "• HR reports normal attendance, but modules tell conflicting stories."
tf.add_paragraph().text = "• Payroll was paying OT from approved requests rather than actual biometric punches."

# Slide 3: Solution Architecture
slide = prs.slides.add_slide(slide_layout)
slide.shapes.title.text = "Our Pipeline Solution"
tf = slide.placeholders[1].text_frame
tf.text = "• Stage 1: Ingested 12 raw CSV/JSON files untouched with hashing."
tf.add_paragraph().text = "• Stage 2 & 3: Cleaned employees, fixed device clock drifts, applied 2-minute rule."
tf.add_paragraph().text = "• Stage 4 & 5: Built Star Schema and answered 10 CFO queries via SQL."

# Slide 4: Finding 1 - Overtime Inflation
slide = prs.slides.add_slide(slide_layout)
slide.shapes.title.text = "CFO Finding 1: Unverified OT Payouts"
tf = slide.placeholders[1].text_frame
tf.text = "• Over 18% of paid overtime occurred on days with zero actual overtime worked."
tf.add_paragraph().text = "• Root Cause: Managers approved OT requests blindly."
tf.add_paragraph().text = "• Action: Link payroll payouts strictly to `fact_attendance_daily` biometric calculations."

# Slide 5: Finding 2 & 3 - Clock Drift & Separation Leakage
slide = prs.slides.add_slide(slide_layout)
slide.shapes.title.text = "CFO Finding 2 & 3: Hardware & Payroll Leaks"
tf = slide.placeholders[1].text_frame
tf.text = "• Finding 2: Device DEV-04 clock was off by 1 hour, causing false late flags."
tf.add_paragraph().text = "• Solution: Enforce NTP synchronization across all terminals."
tf.add_paragraph().text = "• Finding 3: 4 terminated employees received post-separation pay."
tf.add_paragraph().text = "• Solution: Automate stop-pay locks upon SCD Type 2 status change."

# Save presentation
prs.save("indus_weave_cfo_presentation.pptx")
print("PowerPoint presentation 'indus_weave_cfo_presentation.pptx' generated successfully!")
